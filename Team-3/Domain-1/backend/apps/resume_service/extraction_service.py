"""
SwipeX Resume Extraction Service (Milestone 3 — Day 3)

Integrates Domain-1 Resume Service with Intern 4's (Domain-3) completed
extraction and NLP skill extraction pipeline:
  - Text extraction: extraction.pdf_extractor.extract_pdf_text
  - Skill extraction: skills.skill_extractor.extract_skills
  - DOCX text extraction: python-docx with Intern 4's skill extractor

Follows strict domain boundaries: imports and executes Intern 4's module
without modifying any Domain-3 source code.
"""

import sys
import logging
from pathlib import Path
from typing import Union, List, Tuple
from django.conf import settings

logger = logging.getLogger("resume_service.extraction")

# Resolve Domain-3 Milestone3work directory
DOMAIN_3_PATH = Path(settings.BASE_DIR).parent.parent / "Domain-3" / "Milestone3work"
if str(DOMAIN_3_PATH) not in sys.path and DOMAIN_3_PATH.exists():
    sys.path.insert(0, str(DOMAIN_3_PATH))


def extract_docx_text(file_path: Union[str, Path]) -> str:
    """
    Extract readable text from a DOCX resume file.
    
    Args:
        file_path: Path to .docx file
        
    Returns:
        Cleaned text from paragraphs and tables.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Resume file not found: {path}")

    try:
        import docx
        doc = docx.Document(path)
        lines = []
        for p in doc.paragraphs:
            t = p.text.strip()
            if t:
                lines.append(t)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    ct = cell.text.strip()
                    if ct and ct not in lines:
                        lines.append(ct)
        extracted = "\n".join(lines).strip()
        if not extracted:
            raise ValueError("No readable text found in DOCX file.")
        return extracted
    except Exception as e:
        if isinstance(e, (FileNotFoundError, ValueError)):
            raise
        raise ValueError(f"Could not read DOCX resume: {e}")


def extract_resume_content(file_path: Union[str, Path]) -> Tuple[str, List[str]]:
    """
    Extract readable text and skills from a stored resume file using
    Intern 4's completed extraction modules.

    Args:
        file_path: Absolute Path to the resume file on disk (.pdf or .docx).

    Returns:
        tuple (parsed_text: str, extracted_skills: list[str])

    Raises:
        FileNotFoundError: If the file does not exist on disk.
        ValueError: If file is unsupported, empty, or unreadable.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Resume file not found on disk: {path}")

    ext = path.suffix.lower()

    # Ensure Domain-3 is in sys.path
    if str(DOMAIN_3_PATH) not in sys.path and DOMAIN_3_PATH.exists():
        sys.path.insert(0, str(DOMAIN_3_PATH))

    # Step 1: Text extraction
    if ext == ".pdf":
        try:
            from extraction.pdf_extractor import extract_pdf_text
            parsed_text = extract_pdf_text(path)
        except ImportError as e:
            logger.error(f"Failed to import Intern 4 pdf_extractor: {e}")
            raise RuntimeError(f"Downstream extraction module unavailable: {e}")
        except Exception as e:
            logger.warning(f"PDF extraction error for {path}: {e}")
            raise ValueError(f"Could not parse PDF resume: {e}")
    elif ext == ".docx":
        parsed_text = extract_docx_text(path)
    else:
        raise ValueError(f"Unsupported resume file extension '{ext}'. Accepted: .pdf, .docx")

    # Step 2: Skill extraction using Intern 4's spaCy PhraseMatcher
    try:
        from skills.skill_extractor import extract_skills
        extracted_skills = extract_skills(parsed_text)
    except ImportError as e:
        logger.error(f"Failed to import Intern 4 skill_extractor: {e}")
        raise RuntimeError(f"Downstream skill extraction module unavailable: {e}")
    except Exception as e:
        logger.warning(f"Skill extraction error: {e}")
        # Fallback to empty skills if skill extraction fails on valid text
        extracted_skills = []

    logger.info(
        f"Extraction complete for {path.name}: {len(parsed_text)} chars, "
        f"{len(extracted_skills)} skills ({', '.join(extracted_skills[:5])}...)"
    )

    return parsed_text, extracted_skills
